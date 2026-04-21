"""
Generate marketing-genius.pptx from the 10-slide HTML presentation.
Run: python3 scripts/make_pptx.py
Output: marketing-genius.pptx
"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx_helpers import (
    rect, textbox, multiline_textbox, label,
    add_header, add_footer, add_faq_tag, add_answer_box, add_bullet_card,
    left_bar,
    W, H, HW, PAD,
    BLUE, DARK, LIGHT, WHITE, ACCENT, TMID, RED_ERR, RED_BG,
)

# ── Presentation setup ──────────────────────────────────────
prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]   # completely blank layout

HDR_H = Inches(0.55)
FTR_H = Inches(0.38)
BODY_Y = HDR_H
BODY_H = H - HDR_H - FTR_H
MID_Y  = BODY_Y + BODY_H / 2   # vertical midpoint of body

# ── Slide 01: Title split ───────────────────────────────────
sl = prs.slides.add_slide(BLANK)

# Left panel (light)
rect(sl, 0, 0, HW, H, fill=LIGHT)
# Right panel (blue)
rect(sl, HW, 0, HW, H, fill=BLUE)
# Left accent column
rect(sl, 0, 0, Inches(0.06), H, fill=BLUE)

# LEFT content
lx = Inches(0.6)
ly = Inches(1.8)
lw = HW - Inches(1.0)

label(sl, '2026 병원 성장 전략 가이드', lx, ly, lw,
      size=9, color=BLUE)
multiline_textbox(sl, [
    {'text': "왜 잘 나가는 원장님들은", 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': "'마케팅천재'를 찾을까요?", 'size': 22, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
], lx, ly + Inches(0.45), lw, Inches(1.4))
rect(sl, lx, ly + Inches(1.95), Inches(1.0), Inches(0.05), fill=BLUE)
textbox(sl, "수년 간 1,000개+ 병원 마케팅 데이터를 분석한 결과, 답은 명확했습니다.",
        lx, ly + Inches(2.1), lw, Inches(0.7),
        size=11, color=TMID, wrap=True, font='Noto Sans KR')

# Contact box
cb_y = ly + Inches(3.0)
rect(sl, lx, cb_y, Inches(0.05), Inches(0.75), fill=BLUE)
rect(sl, lx + Inches(0.05), cb_y, lw - Inches(0.05), Inches(0.75),
     fill=RGBColor(0xeb, 0xf2, 0xfc))
textbox(sl, '마케팅천재 · 대표 정귀명', lx + Inches(0.2), cb_y + Inches(0.08),
        lw - Inches(0.3), Inches(0.32), size=10, bold=True, color=DARK, font='Noto Sans KR')
textbox(sl, '010-5848-9572 (카톡/전화)', lx + Inches(0.2), cb_y + Inches(0.38),
        lw - Inches(0.3), Inches(0.3), size=9, color=TMID, font='Noto Sans KR')

# RIGHT content
rx = HW + Inches(0.7)
rw = HW - Inches(1.2)
ry = Inches(2.0)

textbox(sl, '1,000+', rx, ry, rw, Inches(1.2),
        size=58, bold=True, color=WHITE, font='Noto Serif KR')
textbox(sl, '병원 마케팅 데이터 분석', rx, ry + Inches(1.15), rw, Inches(0.35),
        size=10, color=RGBColor(0xff,0xff,0xff), font='Noto Sans KR')

bullets = [
    '브랜드 블로그 × 플레이스 시너지',
    '신환 · 내방 폭발적 증가',
    '마케팅 ROAS 평균 1,000%+',
]
for i, b in enumerate(bullets):
    by = ry + Inches(1.7) + i * Inches(0.45)
    rect(sl, rx, by + Inches(0.09), Inches(0.08), Inches(0.08),
         fill=ACCENT)
    textbox(sl, b, rx + Inches(0.18), by, rw - Inches(0.2), Inches(0.38),
            size=10, color=RGBColor(0xff,0xff,0xff), font='Noto Sans KR')

# Slide number
textbox(sl, '01 / 10', W - Inches(1.6), H - Inches(0.5), Inches(1.5), Inches(0.4),
        size=8, color=RGBColor(0xff,0xff,0xff), align=PP_ALIGN.RIGHT, font='Noto Sans KR')

# ── Slide 02: 핵심 결론 (dark) ──────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=DARK)
add_header(sl, '핵심 결론')
add_footer(sl, '02 / 10')

cx = PAD
cy = BODY_Y + Inches(0.3)

textbox(sl, '"', cx, cy, Inches(1.0), Inches(1.0),
        size=60, bold=True, color=BLUE, font='Noto Serif KR')
multiline_textbox(sl, [
    {'text': '병원 마케팅에서', 'size': 20, 'bold': True, 'color': WHITE, 'font': 'Noto Serif KR'},
    {'text': '[브랜드 블로그]와 [플레이스] 조합을', 'size': 20, 'bold': True, 'color': ACCENT, 'font': 'Noto Serif KR'},
    {'text': '이길 수 있는 매체는 없습니다.', 'size': 20, 'bold': True, 'color': WHITE, 'font': 'Noto Serif KR'},
], cx, cy + Inches(0.8), W - PAD * 2, Inches(1.6))

textbox(sl, '이 두 고리가 맞물렸을 때 \'신환\'과 \'내방\'이라는 결과가 폭발적으로 터져 나왔습니다.',
        cx, cy + Inches(2.5), W - PAD * 2, Inches(0.6),
        size=11, color=RGBColor(0xcc,0xcc,0xcc), wrap=True, font='Noto Sans KR')

stats = [('50x', '파워링크 대비 효율'), ('1,000%', '평균 ROAS'), ('6개월', 'J커브 점유 타임라인')]
for i, (num, lbl) in enumerate(stats):
    sx = cx + i * Inches(3.5)
    sy = cy + Inches(3.2)
    textbox(sl, num, sx, sy, Inches(3.0), Inches(0.8),
            size=30, bold=True, color=ACCENT, font='Noto Serif KR')
    textbox(sl, lbl, sx, sy + Inches(0.75), Inches(3.0), Inches(0.35),
            size=9, color=RGBColor(0x99,0x99,0x99), font='Noto Sans KR')

# ── Slide 03: FAQ 01 ────────────────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=WHITE)
add_header(sl, 'FAQ 01')
add_footer(sl, '03 / 10')

cx = PAD; cy = BODY_Y + Inches(0.25); cw = W - PAD * 2

add_faq_tag(sl, 'FAQ 01', cx, cy)
multiline_textbox(sl, [
    {'text': '요즘 시대에 블로그를', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '꼭 해야 하나요?', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
], cx, cy + Inches(0.4), cw, Inches(1.2))
add_answer_box(sl, '네, 2026년 현재까지도 가장 안정적으로 내원을 만드는 마케팅은 단연코 블로그입니다.',
               cx, cy + Inches(1.7), cw, Inches(0.72))

textbox(sl, '"가장 먼저, 그리고 반드시 브랜드블로그부터 시작하셔야 합니다."',
        cx, cy + Inches(2.55), cw, Inches(0.4),
        size=11, color=TMID, font='Noto Sans KR')

cards = [
    ('⚖️', '까다로운 의료광고법', '의료인 명의의 브랜드 블로그는 원장님의 진심과 전문성을 전달할 수 있는 유일한 공식 채널'),
    ('📈', '복리로 쌓이는 자산', '매달 20~30개씩 포스팅이 누적되어 검색 상단을 장악하는 \'디지털 부동산\'이 됩니다'),
    ('🤖', 'AI 검색이 선택하는 데이터', '2026년 네이버 AI 답변 시스템은 \'출처가 명확하고 깊이 있는 정보\'를 우선순위로 선택'),
]
card_w = (cw - Inches(0.3)) / 3
for i, (ic, ti, de) in enumerate(cards):
    add_bullet_card(sl,
                    cx + i * (card_w + Inches(0.15)),
                    cy + Inches(3.05),
                    card_w, Inches(1.4),
                    ti, de, icon=ic)

# ── Slide 04: FAQ 02 ────────────────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=LIGHT)
add_header(sl, 'FAQ 02')
add_footer(sl, '04 / 10')

cx = PAD; cy = BODY_Y + Inches(0.25); cw = W - PAD * 2

add_faq_tag(sl, 'FAQ 02', cx, cy)
multiline_textbox(sl, [
    {'text': '블로그 해봤는데,', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '효과가 없었어요.', 'size': 22, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
], cx, cy + Inches(0.4), cw, Inches(1.2))
add_answer_box(sl, '효과가 없었던 이유는 환자의 내방을 고려하지 않은 \'편법/저품질 방식\' 때문입니다.',
               cx, cy + Inches(1.7), cw, Inches(0.65))

# stat box
rect(sl, cx, cy + Inches(2.5), Inches(2.0), Inches(1.2), fill=BLUE)
textbox(sl, '50배', cx + Inches(0.15), cy + Inches(2.55), Inches(1.7), Inches(0.7),
        size=36, bold=True, color=WHITE, font='Noto Serif KR')
textbox(sl, '파워링크 대비 효율', cx + Inches(0.15), cy + Inches(3.1), Inches(1.7), Inches(0.35),
        size=9, color=RGBColor(0xcc,0xdd,0xff), font='Noto Sans KR')

# explanation
ex = cx + Inches(2.2)
ew = cw - Inches(2.2)
textbox(sl, '노출궤도에 올라간 블로그는 일 12~18만원 예산으로 하루 2~300명의 방문자를 만들어 냅니다.',
        ex, cy + Inches(2.5), ew, Inches(0.5),
        size=10, color=DARK, wrap=True, font='Noto Sans KR')
textbox(sl, '파워링크로 환산하면 동일량의 방문자를 만드는 데 하루 600만~900만원이 소요됩니다.',
        ex, cy + Inches(3.05), ew, Inches(0.5),
        size=10, color=DARK, wrap=True, font='Noto Sans KR')

# warning
wy = cy + Inches(3.75)
rect(sl, cx, wy, Inches(0.06), Inches(0.72), fill=RED_ERR)
rect(sl, cx + Inches(0.06), wy, cw - Inches(0.06), Inches(0.72), fill=RED_BG)
textbox(sl, "⚠ '싸구려 대행'에 속지 마세요 — AI가 영혼 없이 돌린 원고나, 기계적 배포 방식은 이제 네이버에서 가장 먼저 걸러집니다.",
        cx + Inches(0.2), wy + Inches(0.12), cw - Inches(0.35), Inches(0.55),
        size=9, bold=True, color=RED_ERR, wrap=True, font='Noto Sans KR')

# ── Slide 05: 차별점 ────────────────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=WHITE)
add_header(sl, '차별점')
add_footer(sl, '05 / 10')

cx = PAD; cy = BODY_Y + Inches(0.25); cw = W - PAD * 2

multiline_textbox(sl, [
    {'text': "'마케팅천재'는", 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '단순한 글쓰기 기계가 아닙니다', 'size': 22, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
], cx, cy, cw, Inches(1.2))

pillars = [
    ('01', '상위노출 로직', '네이버 검색 알고리즘을 분석한 전략으로 검색 결과 최상단을 점령합니다'),
    ('02', '검색 예상 키워드 발굴', '환자가 실제로 검색하는 황금 키워드를 발굴해 콘텐츠를 전략적으로 설계합니다'),
    ('03', '심리 기획 원고로 내방 설계', '환자의 심리 흐름을 설계해 블로그를 읽는 것만으로 \'방문 결심\'을 이끌어냅니다'),
]
ph = Inches(1.1)
for i, (num, ti, de) in enumerate(pillars):
    py = cy + Inches(1.35) + i * (ph + Inches(0.12))
    rect(sl, cx, py, Inches(0.06), ph, fill=BLUE)
    rect(sl, cx + Inches(0.06), py, cw - Inches(0.06), ph,
         fill=RGBColor(0xeb, 0xf2, 0xfc))
    textbox(sl, num, cx + Inches(0.2), py + Inches(0.05), Inches(0.6), ph - Inches(0.1),
            size=26, bold=True, color=RGBColor(0xb0,0xc8,0xee), font='Noto Serif KR')
    textbox(sl, ti, cx + Inches(0.95), py + Inches(0.08), cw - Inches(1.1), Inches(0.35),
            size=12, bold=True, color=DARK, font='Noto Sans KR')
    textbox(sl, de, cx + Inches(0.95), py + Inches(0.42), cw - Inches(1.1), Inches(0.52),
            size=10, color=TMID, wrap=True, font='Noto Sans KR')

textbox(sl, '이것만이 2026년 의료 시장에서 가장 안전하게, 확실하게 원장님의 진료실 문을 열게 만드는 방법입니다.',
        cx, cy + Inches(4.75), cw, Inches(0.42),
        size=9, color=TMID, wrap=True, font='Noto Sans KR')

# ── Slide 06: FAQ 03 ────────────────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=LIGHT)
add_header(sl, 'FAQ 03')
add_footer(sl, '06 / 10')

cx = PAD; cy = BODY_Y + Inches(0.25); cw = W - PAD * 2

add_faq_tag(sl, 'FAQ 03', cx, cy)
multiline_textbox(sl, [
    {'text': '시작하면 신환이 바로 늘까요?', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '언제쯤 성과가 보이나요?', 'size': 22, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
], cx, cy + Inches(0.4), cw, Inches(1.2))
add_answer_box(sl, "첫 달부터 환자가 줄 서는 기적은 없습니다. '임계점'을 넘는 순간 반드시 폭발합니다.",
               cx, cy + Inches(1.7), cw, Inches(0.65))
textbox(sl, '마케팅천재는 [플레이스 유입 + 블로그 설득] 시너지로 빌드업 속도를 압도적으로 앞당깁니다.',
        cx, cy + Inches(2.5), cw, Inches(0.45),
        size=11, color=DARK, wrap=True, font='Noto Sans KR')

s6stats = [('10배', '마케팅 투자비용 대비'), ('1,000%+', '평균 ROAS'), ('6개월', 'J커브 점유 시점')]
for i, (num, lbl) in enumerate(s6stats):
    sx = cx + i * Inches(3.8)
    sy = cy + Inches(3.1)
    textbox(sl, num, sx, sy, Inches(3.5), Inches(0.85),
            size=34, bold=True, color=BLUE, font='Noto Serif KR')
    textbox(sl, lbl, sx, sy + Inches(0.85), Inches(3.5), Inches(0.35),
            size=9, color=TMID, font='Noto Sans KR')

# ── Slide 07: 성과 타임라인 ─────────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=WHITE)
add_header(sl, '🚀 성과 타임라인')
add_footer(sl, '07 / 10')

cx = PAD; cy = BODY_Y + Inches(0.25); cw = W - PAD * 2

multiline_textbox(sl, [
    {'text': "데이터가 증명하는 '성과 가속'", 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '타임라인', 'size': 22, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
], cx, cy, cw, Inches(1.2))

timeline = [
    ('1단계 · 1~2개월', '기반 구축기',
     ['병원 정체성 확립, 네이버 AI 신뢰 데이터 축적', '플레이스 상위노출 최적화 동시 진행', '브랜드 블로그 육성 시기'],
     RGBColor(0x99,0xb8,0xe8)),
    ('2단계 · 3~4개월', '확신 유입기',
     ['"블로그 보고 왔어요", "플레이스 보고 전화드렸어요"', '철학에 설득된 \'충성 신환\' 유입 시작'],
     RGBColor(0x55,0x8a,0xd6)),
    ('3단계 · 6개월 이후', 'J커브 점유기',
     ['\'디지털 부동산\' 복리 효과 폭발', '주요 황금 키워드 점유, 경쟁 병원과 격차 확대'],
     BLUE),
]
col_w = (cw - Inches(0.3)) / 3
for i, (sub, title, items, tc) in enumerate(timeline):
    cx2 = cx + i * (col_w + Inches(0.15))
    cy2 = cy + Inches(1.35)
    ch = BODY_H - Inches(1.6)
    rect(sl, cx2, cy2, col_w, ch, fill=LIGHT)
    rect(sl, cx2, cy2, col_w, Inches(0.07), fill=tc)
    textbox(sl, sub, cx2 + Inches(0.12), cy2 + Inches(0.12), col_w - Inches(0.2), Inches(0.3),
            size=9, bold=True, color=BLUE, font='Noto Sans KR')
    textbox(sl, title, cx2 + Inches(0.12), cy2 + Inches(0.45), col_w - Inches(0.2), Inches(0.38),
            size=13, bold=True, color=DARK, font='Noto Serif KR')
    for j, item in enumerate(items):
        iy = cy2 + Inches(0.9) + j * Inches(0.58)
        textbox(sl, '· ' + item, cx2 + Inches(0.12), iy, col_w - Inches(0.25), Inches(0.55),
                size=9, color=TMID, wrap=True, font='Noto Sans KR')

# ── Slide 08: FAQ 04 ────────────────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=LIGHT)
add_header(sl, 'FAQ 04')
add_footer(sl, '08 / 10')

cx = PAD; cy = BODY_Y + Inches(0.25); cw = W - PAD * 2

add_faq_tag(sl, 'FAQ 04', cx, cy)
multiline_textbox(sl, [
    {'text': '우리 병원의 전문성을', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '제대로 담아낼 수 있을까요?', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
], cx, cy + Inches(0.4), cw, Inches(1.2))
add_answer_box(sl, '원장님의 진료 철학과 병원의 핵심 강점을 써내는 것이 저희의 실력입니다.',
               cx, cy + Inches(1.7), cw, Inches(0.65))

s8cards = [
    ('심층 인터뷰', '진료 철학/핵심 강점을 단 한 번의 인터뷰로 압축'),
    ('의료 전문 작가 전담', '전문 지식을 환자 눈높이의 \'설득 언어\'로 재탄생'),
    ('클린 IP & 보안 관리', '네이버가 신뢰하는 환경에서 안전하게 발행'),
]
card_w = (cw - Inches(0.3)) / 3
for i, (ti, de) in enumerate(s8cards):
    add_bullet_card(sl,
                    cx + i * (card_w + Inches(0.15)),
                    cy + Inches(2.55), card_w, Inches(1.35), ti, de)

textbox(sl, '가장 비효율적인 병원은 원장님이 마케팅 원고를 고민하느라 진료와 경영에 에너지를 못 쏟는 것입니다.',
        cx, cy + Inches(4.1), cw, Inches(0.42),
        size=10, color=TMID, wrap=True, font='Noto Sans KR')

# ── Slide 09: FAQ 05 ────────────────────────────────────────
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=WHITE)
add_header(sl, 'FAQ 05')
add_footer(sl, '09 / 10')

cx = PAD; cy = BODY_Y + Inches(0.25); cw = W - PAD * 2

add_faq_tag(sl, 'FAQ 05', cx, cy)
multiline_textbox(sl, [
    {'text': '주변 병원이 이미 선점했는데,', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '승산이 있을까요?', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
], cx, cy + Inches(0.4), cw, Inches(1.2))
add_answer_box(sl, '네, 충분합니다. 결국 복리가 쌓이는 시간 싸움입니다.',
               cx, cy + Inches(1.7), cw, Inches(0.65))

# strategy box
sb_y = cy + Inches(2.55)
sb_h = Inches(2.0)
rect(sl, cx, sb_y, Inches(0.06), sb_h, fill=BLUE)
rect(sl, cx + Inches(0.06), sb_y, cw - Inches(0.06), sb_h, fill=LIGHT)
textbox(sl, "후발주자를 단숨에 1위로 만드는 '최다 점유' 전략",
        cx + Inches(0.2), sb_y + Inches(0.12), cw - Inches(0.35), Inches(0.38),
        size=12, bold=True, color=DARK, font='Noto Sans KR')
textbox(sl, "네이버 허용 계정 '3개'를 300% 활용: 원장님 명의로 성격이 다른 3개 채널을 동시 가동",
        cx + Inches(0.2), sb_y + Inches(0.58), cw - Inches(0.35), Inches(0.38),
        size=10, color=DARK, wrap=True, font='Noto Sans KR')
textbox(sl, "(공식/철학형 · 정보/키워드형 · 스토리형)",
        cx + Inches(0.2), sb_y + Inches(0.98), cw - Inches(0.35), Inches(0.32),
        size=10, color=TMID, font='Noto Sans KR')
textbox(sl, "→ 경쟁 병원이 1개 채널로 싸울 때, 우리는 3개 채널로 검색 결과를 장악합니다.",
        cx + Inches(0.2), sb_y + Inches(1.35), cw - Inches(0.35), Inches(0.42),
        size=10, bold=True, color=BLUE, wrap=True, font='Noto Sans KR')

textbox(sl, '지금 시작해도 6개월이면 충분히 역전 가능합니다.',
        cx, sb_y + Inches(2.15), cw, Inches(0.38),
        size=11, bold=True, color=DARK, font='Noto Sans KR')

# ── Slide 10: CTA split ─────────────────────────────────────
sl = prs.slides.add_slide(BLANK)
# Left panel (blue)
rect(sl, 0, 0, HW, H, fill=BLUE)
# Right panel (light)
rect(sl, HW, 0, HW, H, fill=LIGHT)

# LEFT content
lx = Inches(0.7)
lw = HW - Inches(1.1)
textbox(sl, '마케팅 투자 대비 수익률',
        lx, Inches(1.4), lw, Inches(0.38),
        size=9, bold=True, color=RGBColor(0xcc,0xdd,0xff), font='Noto Sans KR')
textbox(sl, '1,000%', lx, Inches(1.8), lw, Inches(1.3),
        size=64, bold=True, color=WHITE, font='Noto Serif KR')
textbox(sl, '마케팅천재와 함께하는 원장님들이\n실제로 경험하고 있는 평균 ROAS입니다.',
        lx, Inches(3.1), lw, Inches(0.72),
        size=10, color=RGBColor(0xcc,0xdd,0xff), wrap=True, font='Noto Sans KR')

l_bullets = ['브랜드 블로그 × 플레이스 시너지', '6개월 J커브 성장 타임라인', '복리로 쌓이는 디지털 자산']
for i, b in enumerate(l_bullets):
    by = Inches(3.95) + i * Inches(0.45)
    rect(sl, lx, by + Inches(0.1), Inches(0.08), Inches(0.08), fill=ACCENT)
    textbox(sl, b, lx + Inches(0.18), by, lw - Inches(0.2), Inches(0.38),
            size=10, color=WHITE, font='Noto Sans KR')

# RIGHT content
rx = HW + Inches(0.7)
rw = HW - Inches(1.1)
textbox(sl, '지금 바로 시작하세요',
        rx, Inches(1.4), rw, Inches(0.38),
        size=9, bold=True, color=BLUE, font='Noto Sans KR')
multiline_textbox(sl, [
    {'text': '잘 나가는 원장님들이', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '선택한 이유,', 'size': 22, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    {'text': '직접 경험해보세요.', 'size': 22, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
], rx, Inches(1.85), rw, Inches(1.8))

# contact card
cc_y = Inches(3.75)
cc_h = Inches(1.5)
rect(sl, rx, cc_y, rw, cc_h, fill=DARK)
textbox(sl, '마케팅천재', rx + Inches(0.2), cc_y + Inches(0.15), rw - Inches(0.3), Inches(0.48),
        size=18, bold=True, color=WHITE, font='Noto Serif KR')
textbox(sl, '대표 정귀명', rx + Inches(0.2), cc_y + Inches(0.6), rw - Inches(0.3), Inches(0.32),
        size=10, color=RGBColor(0xaa,0xaa,0xaa), font='Noto Sans KR')
textbox(sl, '010-5848-9572', rx + Inches(0.2), cc_y + Inches(0.95), rw - Inches(0.3), Inches(0.4),
        size=16, bold=True, color=ACCENT, font='Noto Sans KR')

textbox(sl, '카카오톡 또는 전화로 문의주시면\n무료 상담 후 맞춤 전략을 제안드립니다.',
        rx, cc_y + Inches(1.62), rw, Inches(0.72),
        size=10, color=TMID, wrap=True, font='Noto Sans KR')

# ── Save ────────────────────────────────────────────────────
out = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'marketing-genius.pptx')
prs.save(out)
print(f'Saved: {out}')
