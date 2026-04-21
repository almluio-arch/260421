"""
Helper utilities for building marketing-genius.pptx
Colors and functions shared across all slides.
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

# ── Color palette ──────────────────────────────────────────
BLUE   = RGBColor(0x1a, 0x6f, 0xd8)
BLUE_DK= RGBColor(0x12, 0x54, 0xa8)
DARK   = RGBColor(0x0f, 0x19, 0x23)
LIGHT  = RGBColor(0xf4, 0xf7, 0xfa)
WHITE  = RGBColor(0xff, 0xff, 0xff)
ACCENT = RGBColor(0xf0, 0xc0, 0x40)
TMID   = RGBColor(0x4a, 0x55, 0x68)
RED_ERR= RGBColor(0xb9, 0x1c, 0x1c)
RED_BG = RGBColor(0xfe, 0xe2, 0xe2)

# ── Slide dimensions (13.33" × 7.5" widescreen) ────────────
W  = Inches(13.333)
H  = Inches(7.5)
HW = Inches(6.667)   # half width
PAD= Inches(0.6)     # standard padding

# ── Shape helpers ───────────────────────────────────────────
def rect(slide, x, y, w, h, fill=None, line=None, line_w=Pt(0)):
    """Add a filled rectangle, no border by default."""
    shp = slide.shapes.add_shape(1, int(x), int(y), int(w), int(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = line
        shp.line.width = int(line_w)
    else:
        shp.line.fill.background()
    return shp

def left_bar(slide, x, y, h, color=BLUE, w=Inches(0.05)):
    """Thin vertical accent bar on left side of a box."""
    return rect(slide, x, y, w, h, fill=color)

# ── Text helpers ────────────────────────────────────────────
def _apply_run(run, text, size, bold, color, font_name):
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    if color:
        f.color.rgb = color
    f.name = font_name

def textbox(slide, text, x, y, w, h,
            size=14, bold=False, color=DARK,
            align=PP_ALIGN.LEFT,
            font='Noto Sans KR',
            wrap=True, italic=False,
            line_spacing=None):
    """Add a simple single-run textbox."""
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        from pptx.util import Pt as PPt
        from pptx.oxml.ns import qn
        pPr = p._p.get_or_add_pPr()
        lnSpc = etree.SubElement(pPr, qn('a:lnSpc'))
        spcPct = etree.SubElement(lnSpc, qn('a:spcPct'))
        spcPct.set('val', str(int(line_spacing * 1000)))
    run = p.add_run()
    _apply_run(run, text, size, bold, color, font)
    if italic:
        run.font.italic = True
    return tb

def multiline_textbox(slide, lines, x, y, w, h,
                      size=14, bold=False, color=DARK,
                      align=PP_ALIGN.LEFT,
                      font='Noto Sans KR',
                      line_spacing=1.4):
    """Add a textbox with multiple paragraphs."""
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        if isinstance(line, dict):
            _apply_run(run, line['text'], line.get('size', size),
                       line.get('bold', bold), line.get('color', color),
                       line.get('font', font))
        else:
            _apply_run(run, line, size, bold, color, font)
    return tb

def label(slide, text, x, y, w, h=Inches(0.35),
          size=9, bold=True, color=BLUE,
          align=PP_ALIGN.LEFT, font='Noto Sans KR'):
    """Small uppercase label / eyebrow text."""
    return textbox(slide, text, x, y, w, h,
                   size=size, bold=bold, color=color,
                   align=align, font=font)

# ── Slide chrome ────────────────────────────────────────────
def add_header(slide, tag_text, bg=DARK):
    """Dark header bar with brand + tag."""
    bar_h = Inches(0.55)
    rect(slide, 0, 0, W, bar_h, fill=bg)
    # left accent
    rect(slide, 0, 0, Inches(0.05), bar_h, fill=BLUE)
    textbox(slide, '마케팅천재', PAD, 0, Inches(3), bar_h,
            size=10, bold=True, color=WHITE,
            align=PP_ALIGN.LEFT, font='Noto Sans KR')
    textbox(slide, tag_text, W - Inches(4), 0, Inches(3.4), bar_h,
            size=9, bold=False, color=RGBColor(0xff,0xff,0xff),
            align=PP_ALIGN.RIGHT, font='Noto Sans KR')

def add_footer(slide, num_text=''):
    """Blue footer bar."""
    bar_h = Inches(0.38)
    bar_y = H - bar_h
    rect(slide, 0, bar_y, W, bar_h, fill=BLUE)
    textbox(slide, '마케팅천재 · 대표 정귀명',
            PAD, bar_y, Inches(5), bar_h,
            size=8, color=RGBColor(0xff,0xff,0xff),
            font='Noto Sans KR')
    if num_text:
        textbox(slide, num_text,
                W - Inches(2), bar_y, Inches(1.7), bar_h,
                size=8, color=RGBColor(0xff,0xff,0xff),
                align=PP_ALIGN.RIGHT, font='Noto Sans KR')

def add_faq_tag(slide, text, x, y):
    """Blue-bordered FAQ tag chip."""
    rect(slide, x, y, Inches(1.2), Inches(0.32),
         fill=RGBColor(0xe8, 0xf0, 0xfb),
         line=BLUE, line_w=Pt(1))
    textbox(slide, text, x + Inches(0.08), y, Inches(1.1), Inches(0.32),
            size=8, bold=True, color=BLUE, font='Noto Sans KR')

def add_answer_box(slide, answer_text, x, y, w, h=Inches(0.75)):
    """Blue left-bordered answer block."""
    rect(slide, x, y, Inches(0.06), h, fill=BLUE)
    rect(slide, x + Inches(0.06), y, w - Inches(0.06), h,
         fill=RGBColor(0xeb, 0xf2, 0xfc))
    textbox(slide, '✅ A.', x + Inches(0.15), y + Inches(0.04),
            Inches(0.6), Inches(0.28),
            size=8, bold=True, color=BLUE, font='Noto Sans KR')
    textbox(slide, answer_text,
            x + Inches(0.15), y + Inches(0.28), w - Inches(0.3), h - Inches(0.28),
            size=10, bold=False, color=DARK, wrap=True, font='Noto Sans KR')

def add_bullet_card(slide, x, y, w, h, title, desc, icon='', top_color=BLUE):
    """Card with colored top border, title and description."""
    rect(slide, x, y, w, h, fill=LIGHT)
    rect(slide, x, y, w, Inches(0.06), fill=top_color)
    tx = x + Inches(0.15)
    ty = y + Inches(0.12)
    tw = w - Inches(0.3)
    if icon:
        textbox(slide, icon, tx, ty, tw, Inches(0.32),
                size=14, font='Noto Sans KR')
        ty += Inches(0.35)
    textbox(slide, title, tx, ty, tw, Inches(0.32),
            size=10, bold=True, color=DARK, font='Noto Sans KR')
    textbox(slide, desc, tx, ty + Inches(0.32), tw, h - (ty - y) - Inches(0.32),
            size=9, color=TMID, wrap=True, font='Noto Sans KR')
