"""Slides 09–13"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx_helpers import (
    rect, textbox, multiline_textbox, label,
    add_header, add_footer, add_faq_tag, add_answer_box, add_bullet_card,
    W, H, HW, PAD,
    BLUE, DARK, LIGHT, WHITE, ACCENT, TMID, RED_ERR, RED_BG,
)

TOTAL = '13'
HDR_H = Inches(0.55)
FTR_H = Inches(0.38)
BODY_Y = HDR_H
BODY_H = H - HDR_H - FTR_H


def build_slides_9_13(prs):
    BLANK = prs.slide_layouts[6]

    # ── S09: FAQ 05 경쟁병원 (최다 점유 + 추가 혜택) ──────────
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, fill=WHITE)
    add_header(sl, 'FAQ 05')
    add_footer(sl, '09 / ' + TOTAL)

    cx, cy, cw = PAD, BODY_Y + Inches(0.2), W - PAD * 2
    add_faq_tag(sl, 'FAQ 05', cx, cy)
    multiline_textbox(sl, [
        {'text': '주변 병원이 이미 선점했는데,', 'size': 21, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
        {'text': '승산이 있을까요?',             'size': 21, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
    ], cx, cy + Inches(0.38), cw, Inches(1.1))
    add_answer_box(sl, '네, 충분합니다. 결국 복리가 쌓이는 시간 싸움입니다.',
                   cx, cy + Inches(1.6), cw, Inches(0.58))

    # strategy box
    sb_y, sb_h = cy + Inches(2.3), Inches(1.55)
    rect(sl, cx, sb_y, Inches(0.06), sb_h, fill=BLUE)
    rect(sl, cx + Inches(0.06), sb_y, cw - Inches(0.06), sb_h, fill=LIGHT)
    textbox(sl, "후발주자를 단숨에 1위로 만드는 '최다 점유' 전략",
            cx + Inches(0.2), sb_y + Inches(0.1), cw - Inches(0.35), Inches(0.35),
            size=12, bold=True, color=DARK, font='Noto Sans KR')
    textbox(sl, "네이버 허용 계정 '3개'를 300% 활용: 원장님 명의로 성격이 다른 3개 채널을 동시 가동\n(공식/철학형 · 정보/키워드형 · 스토리형)",
            cx + Inches(0.2), sb_y + Inches(0.5), cw - Inches(0.35), Inches(0.55),
            size=9, color=TMID, wrap=True, font='Noto Sans KR')
    textbox(sl, '→ 경쟁 병원이 1개 채널로 싸울 때, 우리는 3개 채널로 검색 결과를 장악합니다.',
            cx + Inches(0.2), sb_y + Inches(1.1), cw - Inches(0.35), Inches(0.35),
            size=10, bold=True, color=BLUE, wrap=True, font='Noto Sans KR')

    # extra bullets (page 6 content)
    eb_y = cy + Inches(4.0)
    rect(sl, cx, eb_y, Inches(0.06), Inches(0.9), fill=RGBColor(0x55,0x8a,0xd6))
    rect(sl, cx + Inches(0.06), eb_y, cw - Inches(0.06), Inches(0.9), fill=RGBColor(0xf0,0xf5,0xff))
    textbox(sl, '3배 빠른 데이터 축적',
            cx + Inches(0.2), eb_y + Inches(0.05), cw - Inches(0.35), Inches(0.3),
            size=10, bold=True, color=DARK, font='Noto Sans KR')
    textbox(sl, '단일 블로그 대비 양질의 콘텐츠가 더 빠르게 누적됩니다.',
            cx + Inches(0.2), eb_y + Inches(0.32), cw / 2 - Inches(0.2), Inches(0.3),
            size=9, color=TMID, font='Noto Sans KR')
    textbox(sl, '지역 대세감 형성',
            cx + cw / 2, eb_y + Inches(0.05), cw / 2 - Inches(0.1), Inches(0.3),
            size=10, bold=True, color=DARK, font='Noto Sans KR')
    textbox(sl, "다양한 키워드 반복 노출 → '지역에서 핫한 병원' 인식이 빠르게 형성",
            cx + cw / 2, eb_y + Inches(0.32), cw / 2 - Inches(0.1), Inches(0.35),
            size=9, color=TMID, wrap=True, font='Noto Sans KR')

    textbox(sl, '지금 시작해도 6개월이면 충분히 역전 가능합니다.',
            cx, cy + Inches(5.1), cw, Inches(0.38),
            size=11, bold=True, color=DARK, font='Noto Sans KR')

    # ── S10: FAQ 06+07 통합 (좌: 전환공식 / 우: 안심케어) ────
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, fill=WHITE)
    add_header(sl, 'FAQ 06 + 07')
    add_footer(sl, '10 / ' + TOTAL)

    LW = HW - Inches(0.1)   # left panel width
    RX = HW + Inches(0.1)   # right panel start
    RW = W - RX - PAD       # right panel width

    # Left panel background
    rect(sl, 0, HDR_H, HW, BODY_H, fill=LIGHT)

    # FAQ 06
    lx, ly = PAD, BODY_Y + Inches(0.2)
    add_faq_tag(sl, 'FAQ 06', lx, ly)
    multiline_textbox(sl, [
        {'text': '노출은 되는데', 'size': 18, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
        {'text': '왜 환자가 안 올까요?', 'size': 18, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
    ], lx, ly + Inches(0.38), LW - PAD, Inches(1.0))
    add_answer_box(sl, "상위노출은 '입장권'일 뿐. 환자가 내원을 결정하는 건 결국 '설득력'입니다.",
                   lx, ly + Inches(1.52), LW - PAD, Inches(0.65))

    # 전환 공식
    textbox(sl, "노출을 '내원'으로 바꾸는 전환 공식",
            lx, ly + Inches(2.32), LW - PAD, Inches(0.32),
            size=10, bold=True, color=DARK, font='Noto Sans KR')
    conv = [
        ('클릭은 기술, 내원은 심리', '결핍을 분석해 그 결핍을 채우는 [심리 기획 원고] 작성'),
        ('첫인상에서 신뢰 확보', '병원의 격에 맞는 프리미엄 디자인 결합'),
        ('전문성 있는 컨텐츠 제공', 'AI 검색 로직에 맞춘 깊이 있는 정보로 신뢰 형성'),
    ]
    for i, (bold_t, plain_t) in enumerate(conv):
        iy = ly + Inches(2.75) + i * Inches(0.7)
        rect(sl, lx, iy + Inches(0.05), Inches(0.05), Inches(0.55), fill=BLUE)
        textbox(sl, bold_t, lx + Inches(0.15), iy, LW - PAD - Inches(0.2), Inches(0.28),
                size=10, bold=True, color=DARK, font='Noto Sans KR')
        textbox(sl, plain_t, lx + Inches(0.15), iy + Inches(0.28), LW - PAD - Inches(0.2), Inches(0.35),
                size=9, color=TMID, wrap=True, font='Noto Sans KR')

    # Right panel
    add_faq_tag(sl, 'FAQ 07', RX, BODY_Y + Inches(0.2))
    multiline_textbox(sl, [
        {'text': '저품질/노출 이슈가', 'size': 18, 'bold': True, 'color': DARK, 'font': 'Noto Serif KR'},
        {'text': '생기면 어떡하죠?', 'size': 18, 'bold': True, 'color': BLUE, 'font': 'Noto Serif KR'},
    ], RX, BODY_Y + Inches(0.58), RW, Inches(1.0))
    add_answer_box(sl, "정석을 지키면 저품질은 오지 않습니다.\n만약의 사태에도 마케팅 리스크는 '0'입니다.",
                   RX, BODY_Y + Inches(1.7), RW, Inches(0.72))

    textbox(sl, "'안심 케어' 리스크 관리 정책",
            RX, BODY_Y + Inches(2.57), RW, Inches(0.32),
            size=10, bold=True, color=DARK, font='Noto Sans KR')
    risks = [
        ('의료광고 대응 노하우', '보건소 민원 대응까지 실전 경험 기반으로 관리'),
        ('실시간 모니터링 & 즉각 대응', '노출 상태 점검과 방향 수정으로 리스크 최소화'),
        ('무상 복구 시스템', '정책 변화로 이슈 발생 시 준최적화 블로그를 즉시 무상 제공하여 중단 없이 노출'),
    ]
    for i, (bold_t, plain_t) in enumerate(risks):
        iy = BODY_Y + Inches(3.0) + i * Inches(0.78)
        rect(sl, RX, iy + Inches(0.05), Inches(0.05), Inches(0.65), fill=BLUE)
        textbox(sl, bold_t, RX + Inches(0.15), iy, RW - Inches(0.2), Inches(0.28),
                size=10, bold=True, color=DARK, font='Noto Sans KR')
        textbox(sl, plain_t, RX + Inches(0.15), iy + Inches(0.28), RW - Inches(0.2), Inches(0.42),
                size=9, color=TMID, wrap=True, font='Noto Sans KR')

    # ── S11: FAQ 08 브랜드 블로그 가격 ─────────────────────
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, fill=LIGHT)
    add_header(sl, 'FAQ 08 · 비용')
    add_footer(sl, '11 / ' + TOTAL)

    cx, cy, cw = PAD, BODY_Y + Inches(0.2), W - PAD * 2
    add_faq_tag(sl, 'FAQ 08', cx, cy)
    textbox(sl, '비용은 얼마나 들까요?',
            cx, cy + Inches(0.38), cw, Inches(0.55),
            size=22, bold=True, color=DARK, font='Noto Serif KR')

    # intro
    rect(sl, cx, cy + Inches(1.05), Inches(0.05), Inches(0.6), fill=BLUE)
    textbox(sl, "AI 공장형 광고(월 30~50만원)로는 광고비 회수조차 어렵습니다. 마케팅은 '지출'이 아니라 '투자'입니다. 회수 가능한 예산을 제안합니다.",
            cx + Inches(0.15), cy + Inches(1.05), cw - Inches(0.2), Inches(0.6),
            size=10, color=DARK, wrap=True, font='Noto Sans KR')

    # section title
    textbox(sl, '1. 브랜드 블로그',
            cx, cy + Inches(1.82), cw, Inches(0.4),
            size=16, bold=True, color=BLUE, font='Noto Serif KR')
    textbox(sl, '제공: 전문 기획 + 의료 전문 작가 + 프리미엄 디자인 + 의료법 3단계 검수 + 클린IP',
            cx, cy + Inches(2.25), cw, Inches(0.3),
            size=9, color=TMID, font='Noto Sans KR')

    # 3 price cards
    plans = [
        ('BASIC',   '월 16건 (주 4회)',  '100만 원', '브랜드 기초 다지기', False),
        ('PREMIUM', '월 20건 (주 5회)',  '120만 원', '상위노출 & J커브 진입\n월 20건 계정 1~3개 운영이\n효율이 가장 높습니다.', True),
        ('MASTER',  '월 30건 (매일)',    '165만 원', '지역 점유 & 타지역 노출', False),
    ]
    card_w = (cw - Inches(0.4)) / 3
    for i, (name, freq, price, desc, featured) in enumerate(plans):
        cx2 = cx + i * (card_w + Inches(0.2))
        cy2 = cy + Inches(2.65)
        ch = Inches(2.5)
        bg = BLUE if featured else WHITE
        textc = WHITE if featured else DARK
        midc = WHITE if featured else TMID
        rect(sl, cx2, cy2, card_w, ch, fill=bg,
             line=BLUE, line_w=Pt(1.5))

        textbox(sl, name, cx2, cy2 + Inches(0.15), card_w, Inches(0.4),
                size=14, bold=True, color=textc, align=PP_ALIGN.CENTER, font='Noto Sans KR')
        textbox(sl, freq, cx2, cy2 + Inches(0.55), card_w, Inches(0.3),
                size=9, color=midc, align=PP_ALIGN.CENTER, font='Noto Sans KR')
        textbox(sl, price, cx2, cy2 + Inches(0.88), card_w, Inches(0.52),
                size=22, bold=True, color=ACCENT if featured else BLUE,
                align=PP_ALIGN.CENTER, font='Noto Serif KR')
        textbox(sl, desc, cx2 + Inches(0.1), cy2 + Inches(1.45), card_w - Inches(0.2), Inches(0.95),
                size=9, color=midc, align=PP_ALIGN.CENTER, wrap=True, font='Noto Sans KR')

    # ── S12: 플레이스 상위노출 가격 ────────────────────────
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, H, fill=WHITE)
    add_header(sl, 'FAQ 08 · 비용')
    add_footer(sl, '12 / ' + TOTAL)

    cx, cy, cw = PAD, BODY_Y + Inches(0.2), W - PAD * 2

    textbox(sl, '2. 플레이스 상위노출',
            cx, cy, cw, Inches(0.45),
            size=18, bold=True, color=BLUE, font='Noto Serif KR')
    textbox(sl, '환자가 우리 병원을 발견하는 가장 강력한 통로입니다.',
            cx, cy + Inches(0.5), cw, Inches(0.32),
            size=11, color=DARK, font='Noto Sans KR')

    # service box
    sb_y = cy + Inches(0.95)
    rect(sl, cx, sb_y, Inches(0.05), Inches(0.65), fill=BLUE)
    rect(sl, cx + Inches(0.05), sb_y, cw - Inches(0.05), Inches(0.65), fill=LIGHT)
    textbox(sl, '제공: [지역명 + 진료과] 1-5위 · 진료과별 세부 키워드 확장',
            cx + Inches(0.2), sb_y + Inches(0.12), cw - Inches(0.35), Inches(0.38),
            size=10, bold=True, color=DARK, font='Noto Sans KR')

    # bullets
    for i, b in enumerate(['월 150만 원~  (키워드 경쟁도 및 난이도에 따라 상이)',
                            '1-5위, 최소 25일 노출 보장',
                            '순위 이탈 시, A/S 필수 보장']):
        by = sb_y + Inches(0.82) + i * Inches(0.42)
        rect(sl, cx, by + Inches(0.07), Inches(0.08), Inches(0.08), fill=BLUE)
        textbox(sl, b, cx + Inches(0.18), by, cw - Inches(0.25), Inches(0.35),
                size=10, bold=(i == 0), color=DARK, font='Noto Sans KR')

    # synergy tip (yellow bg)
    st_y = sb_y + Inches(2.15)
    rect(sl, cx, st_y, Inches(0.05), Inches(0.65), fill=ACCENT)
    rect(sl, cx + Inches(0.05), st_y, cw - Inches(0.05), Inches(0.65), fill=RGBColor(0xff,0xf8,0xe0))
    textbox(sl, '💡 시너지 팁: 플레이스 단독 진행보다 [브랜드 블로그]와 결합 시 신환 전환율이 300%+ 이상 급증합니다.',
            cx + Inches(0.2), st_y + Inches(0.08), cw - Inches(0.35), Inches(0.5),
            size=10, bold=False, color=DARK, wrap=True, font='Noto Sans KR')

    # recommendation box
    rb_y = st_y + Inches(0.82)
    rect(sl, cx, rb_y, Inches(0.06), Inches(1.0), fill=BLUE)
    rect(sl, cx + Inches(0.06), rb_y, cw - Inches(0.06), Inches(1.0), fill=RGBColor(0xeb,0xf2,0xfc))
    textbox(sl, '✅ 최소 권장 예산: 월 100 ~ 300만 원',
            cx + Inches(0.2), rb_y + Inches(0.08), cw - Inches(0.35), Inches(0.28),
            size=10, bold=True, color=DARK, font='Noto Sans KR')
    textbox(sl, 'ROAS (투자 대비 매출): 파트너 병원 평균 1,000% +   |   필승 조합: [브랜드 블로그] + [플레이스]',
            cx + Inches(0.2), rb_y + Inches(0.38), cw - Inches(0.35), Inches(0.28),
            size=10, color=DARK, wrap=True, font='Noto Sans KR')
    textbox(sl, '마케팅 확장 (홈페이지·인스타·유튜브·카페·바이럴 등) 은 [브랜드 블로그]+[플레이스] 안착 이후입니다.',
            cx + Inches(0.2), rb_y + Inches(0.68), cw - Inches(0.35), Inches(0.28),
            size=9, color=TMID, wrap=True, font='Noto Sans KR')

    # closing quote
    q_y = rb_y + Inches(1.15)
    rect(sl, cx, q_y, Inches(0.06), Inches(0.75), fill=DARK)
    rect(sl, cx + Inches(0.06), q_y, cw - Inches(0.06), Inches(0.75), fill=RGBColor(0xf0,0xf2,0xf5))
    textbox(sl, '"단언컨대, 병원 마케팅의 정답은 [브랜드 블로그]와 [플레이스]의 시너지에 있습니다. 확장은 그 이후입니다."',
            cx + Inches(0.2), q_y + Inches(0.1), cw - Inches(0.35), Inches(0.55),
            size=10, bold=True, color=DARK, wrap=True, font='Noto Sans KR')

    # ── S13: FAQ 09 + CTA ──────────────────────────────────
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, HW, H, fill=BLUE)
    rect(sl, HW, 0, HW, H, fill=LIGHT)

    # LEFT: CTA
    lx, lw = Inches(0.7), HW - Inches(1.1)
    textbox(sl, '마케팅 투자 대비 수익률',
            lx, Inches(1.1), lw, Inches(0.35),
            size=9, bold=True, color=RGBColor(0xcc,0xdd,0xff), font='Noto Sans KR')
    textbox(sl, '1,000%', lx, Inches(1.5), lw, Inches(1.3),
            size=64, bold=True, color=WHITE, font='Noto Serif KR')
    textbox(sl, '마케팅천재와 함께하는 원장님들이\n실제로 경험하고 있는 평균 ROAS입니다.',
            lx, Inches(2.85), lw, Inches(0.72),
            size=10, color=RGBColor(0xcc,0xdd,0xff), wrap=True, font='Noto Sans KR')

    cta_lines = ['브랜드 블로그 × 플레이스 시너지',
                 '6개월 J커브 성장 타임라인',
                 '복리로 쌓이는 디지털 자산']
    for i, b in enumerate(cta_lines):
        by = Inches(3.65) + i * Inches(0.45)
        rect(sl, lx, by + Inches(0.1), Inches(0.08), Inches(0.08), fill=ACCENT)
        textbox(sl, b, lx + Inches(0.18), by, lw - Inches(0.2), Inches(0.38),
                size=10, color=WHITE, font='Noto Sans KR')

    textbox(sl, "'마케팅천재' 문의의 70% 이상은 원장님께서 직접 주십니다.\n이미 검증된 곳에 의구심을 갖는 것은 성장을 늦출 뿐입니다.",
            lx, Inches(5.05), lw, Inches(0.72),
            size=9, color=RGBColor(0xaa,0xcc,0xf0), wrap=True, font='Noto Sans KR')

    # RIGHT: FAQ09 + Contact
    rx = HW + Inches(0.6)
    rw = HW - Inches(1.0)

    add_faq_tag(sl, 'FAQ 09', rx, Inches(0.7))
    textbox(sl, '최소 계약기간이 있나요?',
            rx, Inches(1.1), rw, Inches(0.42),
            size=16, bold=True, color=DARK, font='Noto Serif KR')

    # answer
    rect(sl, rx, Inches(1.6), Inches(0.05), Inches(1.25), fill=BLUE)
    rect(sl, rx + Inches(0.05), Inches(1.6), rw - Inches(0.05), Inches(1.25), fill=RGBColor(0xeb,0xf2,0xfc))
    textbox(sl, '✅  별도의 의무 계약 기간은 없습니다.',
            rx + Inches(0.18), Inches(1.66), rw - Inches(0.28), Inches(0.3),
            size=10, bold=True, color=DARK, font='Noto Sans KR')
    textbox(sl, '다만, 지역 키워드 점유는 평균 1개월 차부터 시작되며, 신규 환자 유입 체감은 2~3개월 차부터 본격화됩니다. 안정적인 성과 궤도를 위해 최소 3개월 운영을 권장드립니다.',
            rx + Inches(0.18), Inches(2.0), rw - Inches(0.28), Inches(0.78),
            size=9, color=TMID, wrap=True, font='Noto Sans KR')

    # closing message
    textbox(sl, '고민은 확신으로 바꾸고,\n지금 바로 실행의 차이를 경험해 보세요.',
            rx, Inches(3.0), rw, Inches(0.72),
            size=12, bold=True, color=DARK, wrap=True, font='Noto Serif KR')

    # contact card
    cc_y = Inches(3.85)
    rect(sl, rx, cc_y, rw, Inches(1.55), fill=DARK)
    textbox(sl, '마케팅천재', rx + Inches(0.2), cc_y + Inches(0.15), rw - Inches(0.3), Inches(0.48),
            size=18, bold=True, color=WHITE, font='Noto Serif KR')
    textbox(sl, '대표 정귀명', rx + Inches(0.2), cc_y + Inches(0.62), rw - Inches(0.3), Inches(0.3),
            size=10, color=RGBColor(0xaa,0xaa,0xaa), font='Noto Sans KR')
    textbox(sl, '010-5848-9572 (카톡/전화)', rx + Inches(0.2), cc_y + Inches(0.98), rw - Inches(0.3), Inches(0.42),
            size=14, bold=True, color=ACCENT, font='Noto Sans KR')

    textbox(sl, '카카오톡 또는 전화로 문의주시면\n무료 상담 후 맞춤 전략을 제안드립니다.',
            rx, cc_y + Inches(1.65), rw, Inches(0.65),
            size=9, color=TMID, wrap=True, font='Noto Sans KR')
